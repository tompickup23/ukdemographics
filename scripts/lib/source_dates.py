"""Dates and periods read from the source files, so no transform types them.

A transform's `lastUpdated`, `release_date` and the period in its source line used to
be typed into the script. They then went stale the moment the script read a newer
file: citizenship.json said "year ending June 2026 release" beside a release date of
21 May 2026, which was the March edition's. Everything here reads the file itself.

Rules:
  * Match by text, never by cell position. A phrase found in more than one form with
    different values raises, because which one is right is a question for a person.
  * Nothing found raises. A transform that cannot say where its date came from stops.
  * Files that carry no date at all (Stat-Xplore cube JSON, Census CSV extracts) get
    held_copy_date(): the date the copy on disk was written, which is when it was
    retrieved. Callers say so in a comment, because it is not a release date.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]
_MONTH_RX = "(" + "|".join(MONTHS) + ")"
_DAY_RX = r"(\d{1,2})(?:st|nd|rd|th)?"
_WEEKDAY_RX = r"(?:(?:Mon|Tues|Wednes|Thurs|Fri|Satur|Sun)day\s+)?"
_ABBR = {m[:3].lower(): i for i, m in enumerate(MONTHS, start=1)}

# The publication-date phrasings seen on cover and title sheets. Each captures day,
# month name and year. They are deliberately narrow: the same sheets also say when
# the NEXT release is due ("Next update: 27 August 2026", "The next quarterly
# release will be published on ..."), and a loose "published ... <date>" would read
# that instead.
PUBLISHED_PATTERNS = [
    # Home Office: "Published: 21 May 2026"
    re.compile(r"\bPublished:\s*" + _DAY_RX + r"\s+" + _MONTH_RX + r"\s+(\d{4})\b"),
    # ONS: "originally published at 9:30 a.m. on 21 May 2026"
    re.compile(r"\boriginally published at [^\n]{0,30}? on " + _DAY_RX + r"\s+" + _MONTH_RX + r"\s+(\d{4})\b"),
    # NHS England: "Publication Date: 31st October 2024" (label and value in two cells)
    re.compile(r"\bPublication Date:\s*" + _DAY_RX + r"\s+" + _MONTH_RX + r"\s+(\d{4})\b", re.IGNORECASE),
    # Ministry of Justice: "This release was published on Thursday 29 January 2026 at 9:30am"
    re.compile(r"\bThis release was published on " + _WEEKDAY_RX + _DAY_RX + r"\s+" + _MONTH_RX + r"\s+(\d{4})\b"),
]

COVER_SHEET_RX = re.compile(r"cover|title", re.IGNORECASE)


def _date(day: str, month: str, year: str) -> dt.date:
    return dt.date(int(year), MONTHS.index(month.capitalize()) + 1, int(day))


def cover_lines(path) -> list[str]:
    """Text of the workbook's cover or title sheet, one string per row.

    The cells of a row are joined with a space, so a label and its value in two cells
    ("Publication Date:" then "31st October 2024") read as one line. The sheet is
    found by name (Cover_sheet, Cover Sheet, Cover, Title); none raises.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"{p} not found")
    suffix = p.suffix.lower()
    rows: list[list] = []
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            names = [n for n in wb.sheetnames if COVER_SHEET_RX.search(n)]
            if not names:
                raise ValueError(f"{p.name}: no cover or title sheet among {wb.sheetnames}")
            rows = [list(r) for r in wb[names[0]].iter_rows(values_only=True)]
        finally:
            wb.close()
    elif suffix == ".ods":
        import pandas as pd

        sheets = pd.ExcelFile(p, engine="odf").sheet_names
        names = [n for n in sheets if COVER_SHEET_RX.search(n)]
        if not names:
            raise ValueError(f"{p.name}: no cover or title sheet among {sheets}")
        df = pd.read_excel(p, engine="odf", sheet_name=names[0], header=None)
        rows = [[None if pd.isna(c) else c for c in r] for r in df.itertuples(index=False)]
    else:
        raise ValueError(f"{p.name}: cover sheets are read from .xlsx or .ods files only")
    lines = []
    for r in rows:
        cells = [" ".join(str(c).split()) for c in r if c not in (None, "")]
        if cells:
            lines.append(" ".join(cells))
    return lines


def cover_match(path, pattern, what: str) -> re.Match:
    """The one match of `pattern` on the cover sheet; raises on none or on disagreement.

    Several matches are allowed when they all capture the same groups (a cover that
    repeats its own title); different values raise.
    """
    rx = re.compile(pattern) if isinstance(pattern, str) else pattern
    found = [m for line in cover_lines(path) for m in rx.finditer(line)]
    if not found:
        raise ValueError(f"{Path(path).name}: no {what} found on the cover sheet (pattern {rx.pattern!r})")
    values = {m.groups() for m in found}
    if len(values) > 1:
        raise ValueError(f"{Path(path).name}: {what} found with different values {sorted(values)}")
    return found[0]


def published_date(path) -> dt.date:
    """Publication date printed on the cover or title sheet of a workbook."""
    found = set()
    for line in cover_lines(path):
        for rx in PUBLISHED_PATTERNS:
            for m in rx.finditer(line):
                found.add(_date(*m.groups()))
    if not found:
        raise ValueError(
            f"{Path(path).name}: no publication date on the cover sheet. Known phrasings: "
            + "; ".join(rx.pattern for rx in PUBLISHED_PATTERNS)
        )
    if len(found) > 1:
        raise ValueError(f"{Path(path).name}: several different publication dates {sorted(found)}")
    return found.pop()


def date_phrase(d: dt.date) -> str:
    """'21 May 2026'."""
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}"


# Home Office immigration statistics workbooks. The cover reads, row by row:
# "Immigration System Statistics", "year ending March 2026", the table title, "Home
# Office", "Published: 21 May 2026".
_HO_PERIOD_RX = re.compile(r"^year ending " + _MONTH_RX + r" (\d{4})$", re.IGNORECASE)
_FILENAME_PERIOD_RX = re.compile(r"-(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)-(\d{4})", re.IGNORECASE)


def home_office_period(path) -> str:
    """'year ending June 2026', read from a Home Office cover sheet.

    Where the file name also carries a period (`-jun-2026`), the two must agree: a
    file renamed by hand, or saved under the wrong edition, stops here.
    """
    m = cover_match(path, _HO_PERIOD_RX, "'year ending <month> <year>' line")
    month, year = m.group(1).capitalize(), m.group(2)
    fm = _FILENAME_PERIOD_RX.search(Path(path).name)
    if fm and (_ABBR[fm.group(1).lower()], fm.group(2)) != (MONTHS.index(month) + 1, year):
        raise ValueError(
            f"{Path(path).name}: the cover says year ending {month} {year} but the file name says "
            f"{fm.group(1)}-{fm.group(2)}"
        )
    return f"year ending {month} {year}"


def held_copy_date(*paths) -> dt.date:
    """UTC date the newest of `paths` was written to disk.

    For sources that carry no release date in the file (Stat-Xplore cube JSON, Census
    CSV extracts), this is the retrieval date of the copy held, which is the most the
    file can tell us. It is not a publication date and callers must not call it one.
    """
    if not paths:
        raise ValueError("held_copy_date needs at least one path")
    stamps = []
    for p in paths:
        p = Path(p)
        if not p.exists():
            raise FileNotFoundError(f"{p} not found")
        stamps.append(p.stat().st_mtime)
    return dt.datetime.fromtimestamp(max(stamps), tz=dt.timezone.utc).date()


def newest(*dates: dt.date) -> str:
    """ISO form of the latest of several input dates, for an output joining sources."""
    if not dates:
        raise ValueError("newest needs at least one date")
    return max(dates).isoformat()


# Stat-Xplore quarter labels: "Oct-25 to Dec-25".
_QTR_LABEL_RX = re.compile(r"^(\w{3})-(\d{2}) to (\w{3})-(\d{2})$")


def _quarter_bounds(label: str):
    m = _QTR_LABEL_RX.match(label.strip())
    if not m:
        raise ValueError(f"unrecognised Stat-Xplore quarter label {label!r}")
    m1, y1, m2, y2 = m.groups()
    try:
        start = (2000 + int(y1), _ABBR[m1.lower()])
        end = (2000 + int(y2), _ABBR[m2.lower()])
    except KeyError:
        raise ValueError(f"unrecognised month in Stat-Xplore quarter label {label!r}") from None
    if end[1] % 3 != 0:
        raise ValueError(f"Stat-Xplore quarter label {label!r} does not end on a calendar quarter")
    return start, end


def quarter_field(cube: dict) -> dict:
    """The quarter dimension of a Stat-Xplore cube, found by its URI or label."""
    hits = [f for f in cube.get("fields", [])
            if "QTR" in f.get("uri", "") or "quarter" in f.get("label", "").lower()]
    if len(hits) != 1:
        raise ValueError(f"expected one quarter field in the cube, found {len(hits)}")
    return hits[0]


def rolling_year(cube: dict, quarters: int = 4) -> dict:
    """The period of the last `quarters` quarters in a Stat-Xplore cube.

    Returns {'phrase': 'rolling year ending Q4 2025', 'months': 'January 2025 to
    December 2025'}, read from the cube's own quarter labels.
    """
    labels = [it["labels"][0] for it in quarter_field(cube)["items"]]
    if len(labels) < quarters:
        raise ValueError(f"the cube holds {len(labels)} quarters, fewer than {quarters}")
    first, _ = _quarter_bounds(labels[-quarters])
    _, last = _quarter_bounds(labels[-1])
    return {
        "phrase": f"rolling year ending Q{last[1] // 3} {last[0]}",
        "months": f"{MONTHS[first[1] - 1]} {first[0]} to {MONTHS[last[1] - 1]} {last[0]}",
    }


def month_label(label: str) -> str:
    """'Jan-26' (a Stat-Xplore month label) as 'January 2026'."""
    m = re.match(r"^(\w{3})-(\d{2})$", label.strip())
    if not m or m.group(1).lower() not in _ABBR:
        raise ValueError(f"unrecognised Stat-Xplore month label {label!r}")
    return f"{MONTHS[_ABBR[m.group(1).lower()] - 1]} {2000 + int(m.group(2))}"
