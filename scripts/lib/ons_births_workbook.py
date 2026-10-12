"""Read the ONS "Parents' country of birth" workbook by text, not by position.

Shared by transform_births_by_mother_cob.py and transform_births_history.py.

Both transforms used to open `2024birthsbyparentscountry.xlsx` by name, take a sheet by
name and read columns by number. The 2025 edition changed the file name (`birth`, not
`births`), every Table 6a header ("Code" became "Area of usual residence Code",
"Number of live births to UK-born mothers" became "Live births to UK-born mothers") and
the order of the sheets. Nothing failed: a positional read would simply have carried on.

So everything here is found by what it says:
  - the file by its stable stem, newest edition first (scripts/lib/newest_release.py);
  - a table by the title in its first cell, which also gives the year it counts;
  - a column by its header text, failing loudly when a header is missing or ambiguous;
  - the release date from the cover sheet's "Date of release" line;
  - the publisher from the table's own "Source:" line.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl

from newest_release import editions, newest_edition

# 2023: 2023birthsbyparentscountryofbirth.xlsx, 2024: 2024birthsbyparentscountry.xlsx,
# 2025: 2025birthbyparentscountry.xlsx. The year leads; the rest has drifted.
FILE_PATTERN = r"(?P<edition>\d{4})births?byparentscountry(?:ofbirth)?\.xlsx"

DATASET_TITLE = "Parents' country of birth"
DATASET_PAGE = (
    "https://www.ons.gov.uk/peoplepopulationandcommunity/birthsdeathsandmarriages/"
    "livebirths/datasets/parentscountryofbirth"
)


def norm(value) -> str:
    """Header text with line breaks, non-breaking spaces and runs of spaces collapsed."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def resolve(raw_dir: Path) -> Path:
    return newest_edition(raw_dir, FILE_PATTERN)


def all_editions(raw_dir: Path) -> list[Path]:
    return editions(raw_dir, FILE_PATTERN)


def open_workbook(path: Path):
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def rows(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def find_table(wb, title_pattern: str, header_has: str | None = None):
    """(worksheet, title, match) for the one sheet whose first cell matches `title_pattern`.

    `header_has` narrows the search to sheets with a header cell containing that text.
    It is needed because titles are not always unique: the 2024 edition titled both
    Table 6a and Table 6b "Live births by country of birth of mother and area of usual
    residence, 2024", although 6b counts births by either parent's country of birth.
    """
    rx = re.compile(title_pattern, re.IGNORECASE)
    hits = []
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), (None,))
        title = norm(first[0] if first else None)
        m = rx.search(title)
        if not m:
            continue
        if header_has is not None:
            cells = [norm(c).lower() for r in ws.iter_rows(min_row=1, max_row=20, values_only=True) for c in r]
            if not any(header_has.lower() in c for c in cells):
                continue
        hits.append((ws, title, m))
    if len(hits) != 1:
        found = ", ".join(f"{ws.title}: {t}" for ws, t, _ in hits) or "none"
        raise RuntimeError(f"expected one table titled like {title_pattern!r}, found {found}")
    return hits[0]


def release_date(wb) -> str:
    """ISO date from the cover sheet's 'Date of release: 27 May 2026' line."""
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            for cell in r:
                # "27 May 2026" in 2025, "1st July 2025" in 2024.
                m = re.search(r"Date of release:\s*(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]+)\s+(\d{4})", norm(cell))
                if m:
                    return datetime.strptime(" ".join(m.groups()), "%d %B %Y").date().isoformat()
    raise RuntimeError("no 'Date of release:' line in the workbook; the cover sheet has changed")


def publisher(table_rows: list[tuple]) -> str:
    for r in table_rows[:10]:
        m = re.match(r"Source:\s*(.+)", norm(r[0] if r else None))
        if m:
            return m.group(1).strip()
    raise RuntimeError("no 'Source:' line above the table")


def header_index(table_rows: list[tuple], first_cell_pattern: str) -> int:
    rx = re.compile(first_cell_pattern, re.IGNORECASE)
    for i, r in enumerate(table_rows[:20]):
        if r and rx.search(norm(r[0])):
            return i
    raise RuntimeError(f"no header row whose first cell matches {first_cell_pattern!r}")


def column(headers: list[str], label: str, predicate) -> int:
    """Index of the one header satisfying `predicate`; raises naming `label` otherwise."""
    hits = [j for j, h in enumerate(headers) if predicate(h.lower())]
    if len(hits) != 1:
        raise RuntimeError(
            f"column '{label}': expected one matching header, found {len(hits)} "
            f"({[headers[j] for j in hits]})"
        )
    return hits[0]


def as_int(value):
    """An integer cell, or None for a symbol such as [x], [z] or [low]."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    s = norm(value).replace(",", "")
    return int(s) if re.fullmatch(r"-?\d+", s) else None


def as_float(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = norm(value)
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None


def citation(publisher_name: str, year: int, released: str, table_title: str) -> str:
    """The source line, built only from what the workbook says plus the dataset's name.

    From the 2025 edition the standalone bulletin is gone; ONS publishes these tables as
    the "Parents' country of birth" dataset alongside "Births in England and Wales: <year>".
    """
    released_text = datetime.fromisoformat(released).strftime("%-d %B %Y")
    return (
        f"{publisher_name}, Births in England and Wales: {year}, dataset "
        f"\"{DATASET_TITLE}\", {year} edition, released {released_text}. {table_title}."
    )
